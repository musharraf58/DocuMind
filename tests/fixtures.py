"""Build small but real PDF and DOCX files in memory, so no binaries are committed."""

import io
import zipfile

import docx
from pypdf import PdfReader, PdfWriter


def _pdf_string(text: str) -> str:
    return "(" + text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ")"


def build_pdf(pages: list[list[str]]) -> bytes:
    """A valid PDF with one page per entry; each entry is that page's lines.

    An empty list gives a page with no text layer, which is what a scanned
    page looks like to a text extractor. Uses the built-in Helvetica font, so
    nothing needs embedding. Latin-1 text only.
    """
    objects: list[str] = []  # object N is objects[N - 1]

    def add(body: str) -> int:
        objects.append(body)
        return len(objects)

    catalog = add("")  # filled in once the page tree exists
    page_tree = add("")
    font = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    page_ids = []
    for lines in pages:
        ops = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
        for line in lines:
            ops += [f"{_pdf_string(line)} Tj", "T*"]
        ops.append("ET")
        stream = "\n".join(ops)
        content = add(f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream")
        page_ids.append(
            add(
                f"<< /Type /Page /Parent {page_tree} 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font} 0 R >> >> /Contents {content} 0 R >>"
            )
        )

    kids = " ".join(f"{p} 0 R" for p in page_ids)
    objects[page_tree - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>"
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {page_tree} 0 R >>"

    # The cross-reference table records the byte offset of every object.
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


def encrypt_pdf(pdf: bytes, user_password: str) -> bytes:
    """Re-save a PDF encrypted with RC4 (works without optional dependencies)."""
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf)))
    writer.encrypt(user_password, owner_password="owner", algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def build_docx(blocks: list[tuple[str, object]], merge_first_row: int = 0) -> bytes:
    """A real DOCX from ("heading", text), ("paragraph", text) and
    ("table", [[cell, ...], ...]) blocks, in order. merge_first_row=N merges the
    first N cells of each table's first row into one."""
    document = docx.Document()
    for kind, value in blocks:
        if kind == "heading":
            document.add_heading(value, level=1)
        elif kind == "paragraph":
            document.add_paragraph(value)
        elif kind == "table":
            table = document.add_table(rows=len(value), cols=len(value[0]))
            for r, row in enumerate(value):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
            if merge_first_row > 1:
                table.cell(0, 0).merge(table.cell(0, merge_first_row - 1))
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def build_zip(entries: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return out.getvalue()
